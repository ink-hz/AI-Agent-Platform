"""Repair the already imported ZenTao stories from the original 2026 Q4 workbook.

Dry run is the default. This tool never creates, deletes, approves, or reassigns
stories. Run inside the platform API container so it can use its ZenTao identity.
"""

from __future__ import annotations

import argparse
import asyncio
import base64
import hashlib
import html
import json
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from app.sdk_workbench.zentao import ZenTaoClient
from openpyxl import load_workbook

SHEET = "2026年Q4上线版本"
HEADERS = ("平台", "任务", "责任人", "版本", "开始时间", "结束时间", "优先级", "风险点", "完成情况", "备注")
SOURCE_SHA = "41df018a029f1d27417bec3008bbe7aab3becd785cc38b312c563ce18ee8370d"
LEGACY_MARKER = re.compile(r"SDK2026PLAN-41df018a-R(\d{4})")
PRIORITY = {"高": 1, "中": 2, "低": 3}
MAX_WORKBOOK_BYTES = 8 * 1024 * 1024
SAFE_FIELDS = ("id", "title", "product", "module", "pri", "status", "reviewer", "assignedTo", "assignedDate", "plan", "openedBy", "deleted")


class RepairConflict(RuntimeError):
    pass


@dataclass(frozen=True)
class Q4Row:
    excel_row: int
    module: str
    title: str
    owner: str
    version: str
    planned_start: str
    planned_end: str
    priority: str
    risk: str
    progress: str
    note: str
    links: tuple[tuple[str, str, str], ...]


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _date(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return _text(value)


def parse_q4_workbook(payload: bytes) -> tuple[Q4Row, ...]:
    if not payload or len(payload) > MAX_WORKBOOK_BYTES:
        raise RepairConflict("q4_workbook_size_invalid")
    try:
        workbook = load_workbook(BytesIO(payload), read_only=False, data_only=True, keep_links=False)
    except Exception as exc:
        raise RepairConflict("q4_workbook_invalid") from exc
    if SHEET not in workbook.sheetnames:
        raise RepairConflict("q4_sheet_missing")
    sheet = workbook[SHEET]
    if tuple(_text(sheet.cell(1, i).value) for i in range(1, 11)) != HEADERS:
        raise RepairConflict("q4_header_invalid")
    module = ""
    rows: list[Q4Row] = []
    for row_number in range(2, sheet.max_row + 1):
        cells = [sheet.cell(row_number, index) for index in range(1, 11)]
        if _text(cells[0].value):
            module = _text(cells[0].value)
        title = _text(cells[1].value)
        if not title:
            continue
        if not module:
            raise RepairConflict(f"q4_module_missing:{row_number}")
        priority = _text(cells[6].value)
        if priority and priority not in PRIORITY:
            raise RepairConflict(f"q4_priority_invalid:{row_number}")
        links: list[tuple[str, str, str]] = []
        for cell in cells:
            if cell.hyperlink is None or not cell.hyperlink.target:
                continue
            target = cell.hyperlink.target
            parsed = urlsplit(target)
            if parsed.scheme not in {"https", "http"} or not parsed.netloc:
                raise RepairConflict(f"q4_link_invalid:{cell.coordinate}")
            links.append((cell.coordinate, _text(cell.value), target))
        rows.append(Q4Row(
            row_number, module, title, _text(cells[2].value), _text(cells[3].value),
            _date(cells[4].value), _date(cells[5].value), priority,
            _text(cells[7].value), _text(cells[8].value), _text(cells[9].value),
            tuple(links),
        ))
    if not rows:
        raise RepairConflict("q4_no_tasks")
    return tuple(rows)


def validate_manifest(payload: bytes, rows: tuple[Q4Row, ...], manifest: dict) -> dict[int, dict]:
    """Pin this one-time repair to the reviewed workbook and original stories."""
    if hashlib.sha256(payload).hexdigest() != manifest.get("workbook_sha256"):
        raise RepairConflict("q4_workbook_sha_mismatch")
    try:
        identities = {int(key): value for key, value in manifest["rows"].items()}
        ids = [int(value["id"]) for value in identities.values()]
        hashes = [digest for value in identities.values()
                  for digest in (value["legacy_spec_sha256"], value["legacy_verify_sha256"])]
    except (KeyError, TypeError, ValueError, AttributeError) as exc:
        raise RepairConflict("q4_manifest_ids_invalid") from exc
    if (set(identities) != {row.excel_row for row in rows} or len(ids) != len(set(ids))
            or any(not isinstance(digest, str) or not re.fullmatch(r"[0-9a-f]{64}", digest) for digest in hashes)):
        raise RepairConflict("q4_manifest_ids_invalid")
    return identities


def _html_text(value: str) -> str:
    return html.escape(value, quote=True).replace("\r\n", "\n").replace("\n", "<br />")


def render_story(row: Q4Row) -> dict[str, str]:
    """Use only workbook facts. Unknown requesters and acceptance remain unknown."""
    source_note = f"{SHEET}｜{' '.join(row.module.split())}｜Excel第{row.excel_row}行"
    if len(source_note) > 255:
        raise RepairConflict(f"q4_source_note_too_long:{row.excel_row}")
    parts = [f"<p>{_html_text(row.title)}</p>"]
    schedule = []
    if row.version and row.version != "-":
        schedule.append(f"计划版本：{_html_text(row.version)}")
    if row.planned_start:
        schedule.append(f"计划开始：{_html_text(row.planned_start)}")
    if row.planned_end:
        schedule.append(f"计划结束：{_html_text(row.planned_end)}")
    if schedule:
        parts.append(f"<p>{'；'.join(schedule)}</p>")
    if row.risk:
        parts.append(f"<p><strong>风险与依赖：</strong>{_html_text(row.risk)}</p>")
    if row.progress:
        parts.append(f"<p><strong>原计划进展：</strong>{_html_text(row.progress)}</p>")
    if row.note:
        parts.append(f"<p><strong>原计划备注：</strong>{_html_text(row.note)}</p>")
    if row.links:
        items = "".join(
            f'<li><a href="{html.escape(url, quote=True)}">{_html_text(label or cell)}</a></li>'
            for cell, label, url in row.links
        )
        parts.append(f"<p><strong>参考资料：</strong></p><ul>{items}</ul>")
    if row.title == "根据客户需要更新版本":
        parts.append("<p>待澄清：具体客户及场景、目标版本、更新范围和验收方式。</p>")
    return {
        "spec": "".join(parts),
        "verify": "<p>原计划未提供可验证的验收标准，待需求提出方和责任人补充。</p>",
        "source": "other",
        "sourceNote": source_note,
    }


def _normalized(value: object) -> str:
    raw = str(value or "")

    def undo_auto_link(match: re.Match[str]) -> str:
        attributes, label = match.group(1), match.group(2)
        href = re.search(r'\bhref="([^"]+)"', attributes)
        if href and html.unescape(href.group(1)) == html.unescape(label):
            return label
        return match.group(0)

    raw = re.sub(r"<a\b([^>]*)>([^<]*)</a>", undo_auto_link, raw, flags=re.IGNORECASE)
    return " ".join(html.unescape(raw).split())


def assert_source_story(story: dict, row: Q4Row, *, product_id: int,
                        identity: dict | None = None) -> str:
    if identity is not None and int(story.get("id", 0)) != int(identity["id"]):
        raise RepairConflict(f"q4_story_id_changed:{row.excel_row}")
    if int(story.get("product", 0)) != product_id or story.get("deleted") in (1, "1", True):
        raise RepairConflict(f"q4_story_identity:{row.excel_row}")
    if _normalized(story.get("title")) != _normalized(row.title):
        raise RepairConflict(f"q4_story_title:{row.excel_row}")
    if story.get("status") != "reviewing":
        raise RepairConflict(f"q4_story_status:{row.excel_row}")
    if row.priority and int(story.get("pri", 0)) != PRIORITY[row.priority]:
        raise RepairConflict(f"q4_story_priority:{row.excel_row}")
    rendered = render_story(row)
    spec = str(story.get("spec") or "")
    marker = LEGACY_MARKER.search(spec)
    if marker and int(marker.group(1)) == row.excel_row and SOURCE_SHA in spec:
        if int(story.get("version", 0)) != 1:
            raise RepairConflict(f"q4_story_human_edit:{row.excel_row}")
        if story.get("sourceNote") not in ("", None, rendered["sourceNote"]):
            raise RepairConflict(f"q4_story_source_changed:{row.excel_row}")
        if (identity is None or hashlib.sha256(spec.encode()).hexdigest()
                != identity["legacy_spec_sha256"]
                or hashlib.sha256(str(story.get("verify") or "").encode()).hexdigest()
                != identity["legacy_verify_sha256"]):
            raise RepairConflict(f"q4_story_human_edit:{row.excel_row}")
        return "source_done" if story.get("sourceNote") == rendered["sourceNote"] else "legacy"
    if (_normalized(spec) == _normalized(rendered["spec"])
            and _normalized(story.get("verify")) == _normalized(rendered["verify"])):
        if story.get("sourceNote") == rendered["sourceNote"]:
            return "done"
        if story.get("sourceNote") in ("", None):
            return "body_done"
    raise RepairConflict(f"q4_story_human_edit:{row.excel_row}")


def protected_fields(story: dict) -> dict:
    return {key: story.get(key) for key in SAFE_FIELDS}


def build_source_note_sql(rows: tuple[Q4Row, ...], stories: dict[int, dict], *, product_id: int,
                          identities: dict[int, dict] | None = None) -> str:
    """Generate one guarded MariaDB transaction for sourceNote only.

    ZenTao OSS 22.6's story PUT API clears assignedTo and rewrites assignedDate
    when updating sourceNote. A narrowly scoped SQL update avoids that API bug.
    """
    statements = []
    for row in rows:
        story = stories[row.excel_row]
        state = assert_source_story(story, row, product_id=product_id,
                                    identity=identities[row.excel_row] if identities else None)
        if state == "done":
            continue
        if state != "body_done":
            raise RepairConflict(f"q4_body_incomplete:{row.excel_row}")
        note64 = base64.b64encode(render_story(row)["sourceNote"].encode()).decode()
        assignee64 = base64.b64encode(str(story.get("assignedTo") or "").encode()).decode()
        statements.append(
            "  UPDATE zt_story SET sourceNote=CONVERT(FROM_BASE64('" + note64 + "') USING utf8mb4)\n"
            f"  WHERE id={int(story['id'])} AND product={product_id} AND status='reviewing' "
            f"AND version={int(story['version'])}\n"
            "    AND sourceNote='' AND assignedTo=CONVERT(FROM_BASE64('" + assignee64 + "') USING utf8mb4);\n"
            f"  IF ROW_COUNT() <> 1 THEN SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT='q4_source_conflict_{row.excel_row}'; END IF;"
        )
    if not statements:
        return ""
    procedure = "sdk_q4_source_notes_20261021"
    return (
        f"DROP PROCEDURE IF EXISTS {procedure};\n"
        "DELIMITER //\n"
        f"CREATE PROCEDURE {procedure}()\nBEGIN\n"
        "  DECLARE EXIT HANDLER FOR SQLEXCEPTION\n"
        "  BEGIN\n    ROLLBACK;\n    RESIGNAL;\n  END;\n"
        "  START TRANSACTION;\n"
        + "\n".join(statements) + "\n"
        "  COMMIT;\n"
        f"  SELECT {len(statements)} AS updated_rows;\n"
        "END//\nDELIMITER ;\n"
        f"CALL {procedure}();\nDROP PROCEDURE {procedure};\n"
    )


def select_rows(rows: tuple[Q4Row, ...], *, only_row: int | None, limit: int | None) -> tuple[Q4Row, ...]:
    if only_row is not None:
        selected = tuple(row for row in rows if row.excel_row == only_row)
        if len(selected) != 1:
            raise RepairConflict(f"q4_row_not_found:{only_row}")
        return selected
    return rows[:limit]


async def change_story_body(client: ZenTaoClient, before: dict, row: Q4Row, *, product_id: int,
                            identity: dict) -> dict:
    state = assert_source_story(before, row, product_id=product_id, identity=identity)
    if state not in {"legacy", "source_done"}:
        raise RepairConflict(f"q4_change_not_needed:{row.excel_row}")
    rendered = render_story(row)
    result = await client._request(
        "POST", f"stories/{before['id']}/change",
        payload={
            "title": html.unescape(before["title"]),
            "spec": rendered["spec"],
            "verify": rendered["verify"],
            "reviewer": before.get("reviewer") or [],
            "lastEditedDate": before.get("lastEditedDate") or "",
            "comment": "依据2026年Q4原始工作表修正导入描述",
        },
        api_version=1,
        expect_envelope=False,
    )
    if str(result.get("id")) != str(before["id"]):
        raise RepairConflict(f"q4_change_response_invalid:{row.excel_row}")
    after = await client.get_story(int(before["id"]))
    if protected_fields(after) != protected_fields(before):
        raise RepairConflict(f"q4_unexpected_side_effect:{row.excel_row}")
    expected_state = "done" if state == "source_done" else "body_done"
    if (assert_source_story(after, row, product_id=product_id, identity=identity) != expected_state
            or int(after.get("version", 0)) != int(before.get("version", 0)) + 1):
        raise RepairConflict(f"q4_readback_failed:{row.excel_row}")
    return after


async def _fetch_stories(client: ZenTaoClient, rows: tuple[Q4Row, ...], product_id: int,
                         identities: dict[int, dict]) -> dict[int, dict]:
    first = await client.list_stories(product_id, 1)
    pager = first.get("pager")
    page_total = pager.get("pageTotal") if isinstance(pager, dict) else None
    if not isinstance(page_total, int) or not 1 <= page_total <= 100:
        raise RepairConflict("q4_pager_invalid")
    summaries = list(first["stories"])
    for page in range(2, page_total + 1):
        summaries.extend((await client.list_stories(product_id, page))["stories"])
    mapped: dict[int, dict] = {}
    for summary in summaries:
        story = await client.get_story(int(summary["id"]))
        marker = LEGACY_MARKER.search(str(story.get("spec") or ""))
        if marker:
            key = int(marker.group(1))
        else:
            source_note = str(story.get("sourceNote") or "")
            match = re.match(rf"^{re.escape(SHEET)}｜.*｜Excel第(\d+)行$", source_note)
            if not match:
                continue
            key = int(match.group(1))
        if key in mapped:
            raise RepairConflict(f"q4_duplicate_story:{key}")
        mapped[key] = story
    expected = {row.excel_row for row in rows}
    if set(mapped) != expected:
        raise RepairConflict(f"q4_coverage_mismatch:missing={sorted(expected-set(mapped))}:extra={sorted(set(mapped)-expected)}")
    for row in rows:
        assert_source_story(mapped[row.excel_row], row, product_id=product_id,
                            identity=identities[row.excel_row])
    return mapped


async def run(workbook: bytes, *, apply: bool, limit: int | None, only_row: int | None,
              output_dir: Path) -> dict:
    manifest = json.loads(Path(__file__).with_name("sdk_q4_manifest.json").read_text())
    rows = parse_q4_workbook(workbook)
    identities = validate_manifest(workbook, rows, manifest)
    if len(rows) != 67 or len({row.module for row in rows}) != 16 or sum(len(row.links) for row in rows) != 8:
        raise RepairConflict("q4_source_shape_mismatch")
    if {identity["id"] for identity in identities.values()} != set(range(3, 70)):
        raise RepairConflict("q4_manifest_ids_invalid")
    product_id = int(os.environ["PLATFORM_SDK_ZENTAO_PRODUCT_ID"])
    client = ZenTaoClient(
        os.environ["PLATFORM_SDK_ZENTAO_BASE_URL"],
        os.environ["PLATFORM_SDK_ZENTAO_ACCOUNT"],
        os.environ["PLATFORM_SDK_ZENTAO_PASSWORD_FILE"],
    )
    try:
        output_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
        os.chmod(output_dir, 0o700)
        snapshot_path = output_dir / "preflight.json"
        if snapshot_path.exists():
            prior = json.loads(snapshot_path.read_text())
            if prior["workbook_sha256"] != hashlib.sha256(workbook).hexdigest():
                raise RepairConflict("q4_snapshot_workbook_changed")
            saved_ids = {int(item["excel_row"]): int(item["old_story"]["id"]) for item in prior["rows"]}
            if set(saved_ids) != {row.excel_row for row in rows}:
                raise RepairConflict("q4_snapshot_coverage_changed")
            if any(saved_ids[row.excel_row] != identities[row.excel_row]["id"] for row in rows):
                raise RepairConflict("q4_snapshot_story_id_changed")
            stories = {row.excel_row: await client.get_story(saved_ids[row.excel_row]) for row in rows}
            for row in rows:
                assert_source_story(stories[row.excel_row], row, product_id=product_id,
                                    identity=identities[row.excel_row])
        else:
            stories = await _fetch_stories(client, rows, product_id, identities)
        snapshot = {
            "sheet": SHEET,
            "workbook_sha256": hashlib.sha256(workbook).hexdigest(),
            "rows": [
                {"excel_row": row.excel_row, "old_story": stories[row.excel_row],
                 "new_fields": render_story(row)}
                for row in rows
            ],
        }
        if not snapshot_path.exists():
            snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2))
            os.chmod(snapshot_path, 0o600)
        result = {"total": len(rows), "already_done": 0, "written": 0, "dry_run": not apply, "rows": []}
        for row in select_rows(rows, only_row=only_row, limit=limit):
            before = await client.get_story(int(stories[row.excel_row]["id"]))
            state = assert_source_story(before, row, product_id=product_id,
                                        identity=identities[row.excel_row])
            if state in {"done", "body_done"}:
                result["already_done"] += 1
                result["rows"].append({"row": row.excel_row, "story": before["id"], "state": state})
                continue
            if not apply:
                result["rows"].append({"row": row.excel_row, "story": before["id"], "state": "preview"})
                continue
            await change_story_body(client, before, row, product_id=product_id,
                                    identity=identities[row.excel_row])
            result["written"] += 1
            result["rows"].append({"row": row.excel_row, "story": before["id"], "state": "written"})
            receipt_path = output_dir / "receipts.json"
            receipt_path.write_text(json.dumps(result, ensure_ascii=False, indent=2))
            os.chmod(receipt_path, 0o600)
        current = {row.excel_row: await client.get_story(int(stories[row.excel_row]["id"])) for row in rows}
        if all(assert_source_story(current[row.excel_row], row, product_id=product_id,
                                   identity=identities[row.excel_row])
               in {"body_done", "done"} for row in rows):
            sql = build_source_note_sql(rows, current, product_id=product_id, identities=identities)
            if sql:
                sql_path = output_dir / "source-notes.sql"
                sql_path.write_text(sql)
                os.chmod(sql_path, 0o600)
        return result
    finally:
        await client.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--only-row", type=int)
    args = parser.parse_args()
    if args.limit is not None and args.limit < 1:
        parser.error("--limit must be positive")
    report = asyncio.run(run(args.workbook.read_bytes(), apply=args.apply,
                             limit=args.limit, only_row=args.only_row, output_dir=args.output_dir))
    print(json.dumps({key: value for key, value in report.items() if key != "rows"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
