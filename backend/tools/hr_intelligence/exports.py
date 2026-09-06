from __future__ import annotations

import json
import re
from collections.abc import Mapping, Sequence
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
_FONT_NAME = "STSong-Light"
_SHEETS = ("原始岗位", "AI分析", "来源覆盖", "证据索引", "成本记录")


def _text(value: object, maximum: int = 32767) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        selected = json.dumps(value, ensure_ascii=False, sort_keys=True)
    else:
        selected = str(value)
    return _CONTROL.sub("", selected)[:maximum]


def _xlsx_value(value: object) -> object:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    selected = _text(value)
    if selected.startswith(_FORMULA_PREFIXES):
        return "'" + selected[:32766]
    return selected


def _rows(value: object) -> list[Mapping[str, object]]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        return []
    return [item for item in value if isinstance(item, Mapping)]


def build_markdown(report: Mapping[str, object]) -> bytes:
    jobs = _rows(report.get("jobs"))
    coverage = _rows(report.get("coverage"))
    analyses = _rows(report.get("analysis"))
    usage = _rows(report.get("usage"))
    evidence = _rows(report.get("evidence"))
    aggregates = report.get("aggregates")
    aggregates = aggregates if isinstance(aggregates, Mapping) else {}
    lines = [
        "# HR 招聘全景情报",
        "",
        f"- Bundle ID：`{_text(report.get('bundle_id'))}`",
        f"- 生成时间：{_text(report.get('generated_at'))}",
        f"- 原始岗位数：{len(jobs)}",
        "",
        "## 来源覆盖",
        "",
    ]
    lines.extend(
        f"- {_text(item.get('company_key'))}：{_text(item.get('state'))}"
        + (
            f"，{_text(item.get('job_count'))} 条岗位快照"
            if item.get("job_count") is not None
            else "，岗位数量未知"
        )
        for item in coverage
    )
    labels = {
        "tracks": {
            "social": "社招",
            "campus": "校招",
            "intern": "实习",
            "unknown": "未分类",
        },
        "job_families": {
            "research_development": "研发",
            "quality": "质量",
            "manufacturing": "制造",
            "supply_chain": "供应链",
            "product": "产品",
            "sales_marketing": "销售与市场",
            "operations": "运营与交付",
            "corporate": "职能",
            "other": "其他",
        },
        "directions": {},
    }
    lines.extend(["", "## 确定性概览", ""])
    for key, title in (
        ("tracks", "招聘类型"),
        ("job_families", "岗位族"),
        ("directions", "技术方向（多标签）"),
    ):
        values = aggregates.get(key)
        if not isinstance(values, Mapping):
            continue
        names = labels[key]
        readable = "、".join(
            f"{names.get(str(name), str(name))}：{count}"
            for name, count in sorted(
                values.items(), key=lambda item: (-int(item[1]), str(item[0]))
            )
            if int(count) > 0
        )
        if readable:
            lines.append(f"- {title}：{readable}")
    lines.extend(
        [
            "",
            "> 技术方向为多标签口径，各方向数量不能直接相加为岗位总数。",
            "",
            "## AI 分析",
            "",
        ]
    )
    if analyses:
        for item in analyses:
            response = item.get("response", item)
            if isinstance(response, Mapping):
                lines.extend(
                    [
                        f"### {_text(item.get('scope_key', item.get('kind', '分析'))) or '分析'}",
                        "",
                        _text(response.get("summary")) or "未提供总结",
                        "",
                        "证据化事实：",
                    ]
                )
                for fact in _rows(response.get("facts")):
                    lines.append(
                        "- "
                        f"{_text(fact.get('text'))} "
                        f"([来源]({_text(fact.get('source_url'))})；"
                        f"SHA-256 `{_text(fact.get('evidence_sha256'))}`)"
                    )
                unknowns = response.get("unknowns")
                if isinstance(unknowns, Sequence) and not isinstance(
                    unknowns, (str, bytes)
                ):
                    lines.extend(["", "未知与限制："])
                    lines.extend(f"- {_text(value)}" for value in unknowns)
                lines.append("")
    else:
        lines.extend(["尚无已接受的 AI 分析。", ""])
    lines.extend(
        [
            "## 数据下载",
            "",
            f"共 {len(jobs)} 条可追溯岗位快照。原始岗位明细请查看 report.xlsx。",
            "",
        ]
    )
    lines.extend(["", "## 证据索引", ""])
    lines.extend(
        f"- `{_text(item.get('sha256'))}`｜{_text(item.get('source_url'))}"
        for item in evidence
    )
    lines.extend(["", "## 分析调用记录", ""])
    if not usage:
        lines.append("尚无分析调用记录。")
    for item in usage:
        lines.append(
            "- "
            f"{_text(item.get('provider')) or '未知 Provider'} / "
            f"{_text(item.get('model')) or '未知模型'}："
            f"输入 Token {_text(item.get('input_tokens')) or '不可用'}，"
            f"输出 Token {_text(item.get('output_tokens')) or '不可用'}，"
            f"费用 {_text(item.get('estimated_cost')) or '不可用'}"
        )
    lines.append("")
    return "\n".join(lines).encode("utf-8")


def _paragraph(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(_text(value)).replace("\n", "<br/>"), style)


def _footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont(_FONT_NAME, 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawCentredString(A4[0] / 2, 10 * mm, f"{document.page}")
    canvas.restoreState()


def build_pdf(report: Mapping[str, object]) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont(_FONT_NAME))
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title="HR 招聘全景情报",
    )
    sample = getSampleStyleSheet()
    title = ParagraphStyle(
        "TitleCN",
        parent=sample["Title"],
        fontName=_FONT_NAME,
        fontSize=20,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#102A43"),
    )
    heading = ParagraphStyle(
        "HeadingCN",
        parent=sample["Heading2"],
        fontName=_FONT_NAME,
        fontSize=14,
        leading=20,
        spaceBefore=8,
        spaceAfter=6,
        textColor=colors.HexColor("#0F766E"),
    )
    body = ParagraphStyle(
        "BodyCN",
        parent=sample["BodyText"],
        fontName=_FONT_NAME,
        fontSize=9,
        leading=14,
        textColor=colors.HexColor("#263238"),
    )
    small = ParagraphStyle(
        "SmallCN",
        parent=body,
        fontSize=7.5,
        leading=11,
        textColor=colors.HexColor("#475569"),
    )
    jobs = _rows(report.get("jobs"))
    aggregates = report.get("aggregates", {})
    aggregates = aggregates if isinstance(aggregates, Mapping) else {}
    story = [
        _paragraph("HR 招聘全景情报", title),
        Spacer(1, 5 * mm),
        _paragraph(f"Bundle ID：{_text(report.get('bundle_id'))}", body),
        _paragraph(f"生成时间：{_text(report.get('generated_at'))}", body),
        _paragraph(f"可追溯岗位快照：{len(jobs)} 条", body),
        Spacer(1, 4 * mm),
        _paragraph("确定性概览", heading),
    ]
    overview_rows = []
    labels = {
        "tracks": "招聘类型",
        "job_families": "岗位族",
        "directions": "技术方向（多标签）",
    }
    for key in ("tracks", "job_families", "directions"):
        values = aggregates.get(key)
        if isinstance(values, Mapping):
            ranked = sorted(
                ((str(name), value) for name, value in values.items()),
                key=lambda item: (-int(item[1]), item[0]),
            )
            overview_rows.append(
                [
                    _paragraph(labels[key], body),
                    _paragraph(
                        "、".join(f"{name} {value}条" for name, value in ranked),
                        body,
                    ),
                ]
            )
    if overview_rows:
        overview = Table(overview_rows, colWidths=(34 * mm, 126 * mm))
        overview.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ECFDF5")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(overview)
    story.extend(
        [
            Spacer(1, 4 * mm),
            _paragraph("来源覆盖", heading),
        ]
    )
    coverage = _rows(report.get("coverage"))
    if coverage:
        data = [
            [
                _paragraph("公司", body),
                _paragraph("状态", body),
                _paragraph("岗位快照", body),
                _paragraph("限制", body),
            ]
        ]
        data.extend(
            [
                _paragraph(item.get("company_key"), body),
                _paragraph(item.get("state"), body),
                _paragraph(item.get("job_count", ""), body),
                _paragraph(item.get("limitations", ""), small),
            ]
            for item in coverage
        )
        table = Table(
            data, colWidths=(42 * mm, 28 * mm, 25 * mm, 65 * mm), repeatRows=1
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DFF5F2")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)
    story.extend(
        [
            Spacer(1, 4 * mm),
            _paragraph("口径与限制", heading),
            _paragraph(
                "公开岗位数不等于HC、预算、录用人数或研发投入。技术方向采用多标签归类，计数会交叉重叠。"
                "当前为单次基线；只有形成跨期可比版本后，才能判断新增、消失和趋势变化。"
                "采集失败或无可解析岗位不代表企业没有招聘。",
                body,
            ),
            PageBreak(),
            _paragraph("AI 分析", heading),
        ]
    )
    analyses = _rows(report.get("analysis"))
    if not analyses:
        story.append(_paragraph("尚无已接受的 AI 分析。", body))
    kind_order = {
        "executive-summary": 0,
        "comparison": 1,
        "company": 2,
        "track": 3,
        "direction": 4,
    }
    for item in sorted(
        analyses,
        key=lambda value: (
            kind_order.get(str(value.get("kind")), 9),
            str(value.get("scope_key", "")),
        ),
    ):
        response = item.get("response", item)
        if isinstance(response, Mapping):
            story.append(
                _paragraph(
                    f"{item.get('scope_key', item.get('kind', '分析'))} "
                    f"｜置信度 {response.get('confidence', '未标注')}",
                    heading,
                )
            )
            story.append(_paragraph(response.get("summary", ""), body))
            inferences = _rows(response.get("inferences"))
            if inferences:
                story.append(_paragraph("关键判断", small))
                for inference in inferences[:3]:
                    story.append(_paragraph(f"— {inference.get('text', '')}", body))
            facts = _rows(response.get("facts"))
            if facts:
                story.append(_paragraph("代表性证据", small))
            for fact in facts[:3]:
                story.append(
                    _paragraph(
                        f"— {fact.get('text', '')}\n来源：{fact.get('source_url', '')}\n"
                        f"SHA-256：{fact.get('evidence_sha256', '')}",
                        small,
                    )
                )
            unknowns = response.get("unknowns")
            if isinstance(unknowns, Sequence) and not isinstance(
                unknowns, (str, bytes)
            ):
                story.append(_paragraph("未知与限制", small))
                for unknown in list(unknowns)[:2]:
                    story.append(_paragraph(f"— {unknown}", small))
    story.extend(
        [
            PageBreak(),
            _paragraph("数据与审计附件", heading),
            _paragraph(
                f"原始岗位明细请查看 report.xlsx 的“原始岗位”工作表，共 {len(jobs)} 条。"
                "AI分析、来源覆盖、证据索引和成本记录分别保存在同一工作簿的独立工作表；"
                "机器可读原文另见 normalized-jobs.jsonl、analysis.json、raw-evidence-index.json 与 checksums.sha256。",
                body,
            ),
            Spacer(1, 4 * mm),
            _paragraph("成本记录", heading),
            _paragraph(
                f"分析单元 {len(_rows(report.get('usage')))} 个。逐单元模型、Token、成本及不可用原因详见 report.xlsx 的“成本记录”工作表。",
                body,
            ),
        ]
    )
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output.getvalue()


def _fit_columns(sheet, maximum: int = 56) -> None:
    for index, column in enumerate(sheet.columns, start=1):
        width = min(
            max((len(_text(cell.value)) for cell in column), default=8) + 2, maximum
        )
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_sheet(
    sheet, rows: list[Mapping[str, object]], headers: tuple[str, ...]
) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append([_xlsx_value(row.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _fit_columns(sheet)


def build_xlsx(report: Mapping[str, object]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    jobs = _rows(report.get("jobs"))
    analysis = _rows(report.get("analysis"))
    coverage = _rows(report.get("coverage"))
    evidence = _rows(report.get("evidence"))
    usage = _rows(report.get("usage"))
    specifications = (
        (
            "原始岗位",
            jobs,
            (
                "company_key",
                "title",
                "location",
                "duty_excerpt",
                "requirement_excerpt",
                "source_url",
                "evidence_sha256",
                "observed_at",
                "status",
            ),
        ),
        (
            "AI分析",
            analysis,
            ("kind", "scope_key", "input_sha256", "response_sha256", "response"),
        ),
        (
            "来源覆盖",
            coverage,
            ("company_key", "state", "observed_at", "job_count", "limitations"),
        ),
        (
            "证据索引",
            evidence,
            ("sha256", "source_url", "mime", "size_bytes", "locator"),
        ),
        (
            "成本记录",
            usage,
            (
                "unit_id",
                "provider",
                "model",
                "input_tokens",
                "output_tokens",
                "estimated_cost",
                "unavailable_reason",
            ),
        ),
    )
    for name, rows, headers in specifications:
        _write_sheet(workbook.create_sheet(name), rows, headers)
    workbook.properties.title = "HR 招聘全景情报"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


__all__ = ["build_markdown", "build_pdf", "build_xlsx"]
