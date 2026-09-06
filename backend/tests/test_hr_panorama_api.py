from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import FastAPI, HTTPException, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.hr.panorama_models import (
    PanoramaReport,
    PublicJobSnapshot,
    PublishedPanorama,
    SourceCollectionAttempt,
    TalentInsightVersion,
    TalentSource,
)
from app.hr.panorama_repository import (
    PanoramaConflict,
    PanoramaNotFound,
    PanoramaUnavailable,
)
from app.hr.panorama_routes import build_panorama_router
from app.hr.panorama_service import PanoramaEvidenceFile

NOW = datetime(2026, 9, 6, 8, tzinfo=UTC)


class FakePanoramaService:
    def __init__(self) -> None:
        owner_id, source_id, batch_id, snapshot_id, observation_id = (
            uuid4(),
            uuid4(),
            uuid4(),
            uuid4(),
            uuid4(),
        )
        source = TalentSource(
            source_id,
            owner_id,
            uuid4(),
            "company",
            "company-union-optech",
            "联合光电",
            ("Union Optech",),
            ("https://example.com/jobs", "https://example.com/campus"),
            True,
            NOW,
            NOW,
        )
        snapshot = PublicJobSnapshot(
            snapshot_id,
            owner_id,
            observation_id,
            None,
            source_id,
            "job-1",
            "高级结构工程师",
            "中山",
            "负责精密结构设计",
            "五年以上量产经验",
            "https://example.com/jobs/1",
            NOW,
            "a" * 64,
            "open",
            NOW,
            batch_id,
            observation_id,
        )
        insight = TalentInsightVersion(
            uuid4(),
            owner_id,
            uuid4(),
            None,
            1,
            (source_id,),
            (snapshot_id,),
            (
                {
                    "fact_id": "f1",
                    "text": "联合光电公开招聘高级结构工程师",
                    "snapshot_id": str(snapshot_id),
                    "observation_id": str(observation_id),
                    "source_url": snapshot.source_url,
                    "observed_at": NOW.isoformat(),
                },
            ),
            ({"text": "结构研发投入明确", "basis_fact_ids": ("f1",)},),
            ({"text": "实际 HC 未公开"},),
            {"结构": 1},
            "结构研发招聘持续",
            None,
            None,
            "hr-intelligence-producer",
            "configured-model-v1",
            NOW,
            batch_id,
        )
        publication = PublishedPanorama(
            uuid4(),
            uuid4(),
            batch_id,
            owner_id,
            insight.insight_version_id,
            "hr",
            "partial",
            (
                {
                    "source_id": str(source_id),
                    "state": "succeeded",
                    "observed_at": NOW.isoformat(),
                    "source_urls": source.approved_urls,
                    "job_count": 1,
                    "channel_failures": {
                        "https://example.com/campus": "source_timeout"
                    },
                },
            ),
            NOW,
        )
        attempt = SourceCollectionAttempt(
            uuid4(),
            batch_id,
            owner_id,
            source_id,
            "https://example.com/jobs",
            1,
            "succeeded",
            None,
            "a" * 64,
            "sha256/aa/" + "a" * 64,
            "application/json",
            len('{"jobs":[{"title":"高级结构工程师"}]}'.encode()),
            1,
            NOW,
            NOW,
        )
        self.report_value = PanoramaReport(
            insight, (source,), (snapshot,), publication, (attempt,)
        )
        self.calls: list[tuple] = []
        self.error: Exception | None = None
        self.empty = False
        self.raw_evidence = '{"jobs":[{"title":"高级结构工程师"}]}'.encode()

    def _result(self, value):
        if self.error is not None:
            raise self.error
        return value

    def current_report(self):
        self.calls.append(("current",))
        return self._result(None if self.empty else self.report_value)

    def list_reports(self, *, limit=100):
        self.calls.append(("list", limit))
        return self._result((self.report_value,))

    def report(self, publication_id):
        self.calls.append(("report", publication_id))
        if publication_id != self.report_value.publication.publication_id:
            raise PanoramaNotFound()
        return self._result(self.report_value)

    def evidence_file(self, publication_id, sha256):
        self.calls.append(("evidence", publication_id, sha256))
        if publication_id != self.report_value.publication.publication_id:
            raise PanoramaNotFound()
        return PanoramaEvidenceFile(
            sha256=sha256,
            mime="application/json",
            body=self.raw_evidence,
        )


def _client(*, entitled: bool = True, stale: bool = False):
    owner_id = uuid4()
    service = FakePanoramaService()
    app = FastAPI()

    async def require_hr_access(_request: Request, *, writable: bool = False):
        assert writable is False
        if not entitled:
            raise HTTPException(403, "HR Agent use denied")
        assert stale in {True, False}
        return owner_id

    app.include_router(build_panorama_router(service, require_hr_access))
    return TestClient(app), service


def test_business_router_exposes_only_published_reads() -> None:
    router = build_panorama_router(FakePanoramaService(), lambda _request, **_: uuid4())
    routes = {
        (route.path, method)
        for route in router.routes
        for method in (getattr(route, "methods", None) or ())
    }

    assert ("/api/hr/panorama/current", "GET") in routes
    assert not any(
        path.startswith("/api/hr/panorama") and method not in {"GET", "HEAD"}
        for path, method in routes
    )
    assert not any("/runs" in path or "/sources" in path for path, _ in routes)


def test_current_returns_204_before_the_first_published_report() -> None:
    client, service = _client()
    service.empty = True

    response = client.get("/api/hr/panorama/current")

    assert response.status_code == 204
    assert response.content == b""


def test_current_keeps_ai_analysis_raw_jobs_and_publication_coverage_separate() -> None:
    client, service = _client()

    response = client.get("/api/hr/panorama/current")

    assert response.status_code == 200
    body = response.json()
    assert body["insight"]["inferences"][0]["text"] == "结构研发投入明确"
    assert body["snapshots"][0]["title"] == "高级结构工程师"
    assert body["snapshots"][0]["content_sha256"] == "a" * 64
    assert body["evidence"][0]["sha256"] == "a" * 64
    assert body["publication"]["coverage_state"] == "partial"
    assert body["publication"]["source_coverage"][0]["channel_failures"] == {
        "https://example.com/campus": "source_timeout"
    }
    assert body["insight"]["run_id"] is None
    assert body["insight"]["production_batch_id"] == str(
        service.report_value.publication.batch_id
    )
    assert response.headers["cache-control"] == "private, no-store"


def test_history_and_detail_use_publication_ids() -> None:
    client, service = _client()
    publication_id = service.report_value.publication.publication_id

    history = client.get("/api/hr/panorama/reports?limit=20")
    detail = client.get(f"/api/hr/panorama/reports/{publication_id}")

    assert history.status_code == detail.status_code == 200
    assert history.json()["items"][0]["publication"]["publication_id"] == str(
        publication_id
    )
    assert "snapshots" not in history.json()["items"][0]
    assert detail.json()["snapshots"][0]["observation_id"]
    assert service.calls == [("list", 20), ("report", publication_id)]


def test_report_exports_preserve_ai_and_original_job_data() -> None:
    client, service = _client()
    base = f"/api/hr/panorama/reports/{service.report_value.publication.publication_id}/export"

    pdf = client.get(f"{base}?format=pdf")
    xlsx = client.get(f"{base}?format=xlsx")

    assert pdf.status_code == xlsx.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    pdf_text = "".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
    workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
    assert "结构研发投入明确" in pdf_text
    assert "高级结构工程师" in pdf_text
    assert {"原始岗位", "AI分析", "来源覆盖", "证据索引"} <= set(
        workbook.sheetnames
    )


def test_archived_raw_source_response_can_be_downloaded_without_mutation() -> None:
    client, service = _client()
    publication_id = service.report_value.publication.publication_id

    response = client.get(
        f"/api/hr/panorama/reports/{publication_id}/evidence/{'a' * 64}"
    )

    assert response.status_code == 200
    assert response.content == service.raw_evidence
    assert response.headers["content-type"] == "application/json"
    assert "attachment" in response.headers["content-disposition"]


def test_excel_neutralizes_formula_cells_from_public_content() -> None:
    client, service = _client()
    snapshot = replace(service.report_value.snapshots[0], title="=HYPERLINK(\"x\")")
    service.report_value = PanoramaReport(
        service.report_value.insight,
        service.report_value.sources,
        (snapshot,),
        service.report_value.publication,
        service.report_value.evidence_attempts,
    )
    publication_id = service.report_value.publication.publication_id

    response = client.get(
        f"/api/hr/panorama/reports/{publication_id}/export?format=xlsx"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=False)

    assert workbook["原始岗位"]["B2"].value.startswith("'=")


@pytest.mark.parametrize(
    ("error", "status", "detail"),
    (
        (PanoramaNotFound("secret"), 404, "HR panorama not found"),
        (PanoramaConflict("secret"), 409, "HR panorama conflict"),
        (PanoramaUnavailable("secret"), 503, "HR panorama unavailable"),
        (TypeError("secret"), 422, "HR panorama request invalid"),
        (ValueError("secret"), 422, "HR panorama request invalid"),
    ),
)
def test_errors_are_sanitized(error, status, detail) -> None:
    client, service = _client()
    service.error = error

    response = client.get("/api/hr/panorama/current")

    assert response.status_code == status
    assert response.json() == {"detail": detail}
    assert "secret" not in response.text
    assert response.headers["cache-control"] == "private, no-store"


def test_unentitled_users_are_blocked_but_stale_accounts_keep_read_access() -> None:
    blocked, blocked_service = _client(entitled=False)
    stale, stale_service = _client(stale=True)

    assert blocked.get("/api/hr/panorama/current").status_code == 403
    assert blocked_service.calls == []
    assert stale.get("/api/hr/panorama/current").status_code == 200
    assert stale_service.calls == [("current",)]


def test_invalid_publication_id_is_rejected_before_service() -> None:
    client, service = _client()
    response = client.get("/api/hr/panorama/reports/not-a-uuid")
    assert response.status_code == 422
    assert service.calls == []
