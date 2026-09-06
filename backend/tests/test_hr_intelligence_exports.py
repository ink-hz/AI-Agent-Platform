from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from tools.hr_intelligence.exports import build_markdown, build_pdf, build_xlsx


def _report_input():
    return {
        "bundle_id": "00000000-0000-4000-8000-000000000001",
        "generated_at": "2026-09-06T08:00:00+00:00",
        "jobs": [
            {
                "company_key": "hesai",
                "title": "算法工程师",
                "location": "上海",
                "source_url": "https://example.com/jobs/1",
                "evidence_sha256": "a" * 64,
            }
        ],
        "coverage": [{"company_key": "hesai", "state": "succeeded", "job_count": 1}],
        "analysis": [],
        "usage": [],
        "aggregates": {"tracks": {"social": 1}},
        "evidence": [{"sha256": "a" * 64, "source_url": "https://example.com/jobs"}],
    }


def test_exports_preserve_raw_jobs_analysis_sources_and_cost_layers() -> None:
    report_input = _report_input()

    markdown = build_markdown(report_input)
    pdf = build_pdf(report_input)
    xlsx = build_xlsx(report_input)
    workbook = load_workbook(BytesIO(xlsx), read_only=True)

    markdown_text = markdown.decode("utf-8")
    assert "原始岗位" in markdown_text
    assert "AI 分析" in markdown_text
    assert "来源覆盖" in markdown_text
    assert pdf.startswith(b"%PDF-")
    assert workbook.sheetnames == [
        "原始岗位",
        "AI分析",
        "来源覆盖",
        "证据索引",
        "成本记录",
    ]
    assert workbook["来源覆盖"]["D2"].value == 1


def test_pdf_keeps_raw_job_detail_in_workbook_instead_of_unbounded_appendix() -> None:
    report_input = _report_input()
    report_input["jobs"] = [
        report_input["jobs"][0] | {"title": f"算法工程师-{index}"}
        for index in range(250)
    ]

    reader = PdfReader(BytesIO(build_pdf(report_input)))
    text = "".join(page.extract_text() or "" for page in reader.pages)

    assert len(reader.pages) <= 4
    assert "原始岗位明细请查看 report.xlsx" in text


def test_human_markdown_uses_readable_summary_instead_of_raw_data_dump() -> None:
    report_input = _report_input()
    report_input["aggregates"] = {
        "tracks": {"social": 1},
        "directions": {"算法": 1},
        "job_families": {"research_development": 1},
    }

    body = build_markdown(report_input).decode("utf-8")

    assert "```json" not in body
    assert "## 原始岗位" not in body
    assert "原始岗位明细请查看 report.xlsx" in body
    assert "社招：1" in body
