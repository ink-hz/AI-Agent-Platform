from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4

import psycopg
import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_control_plane_migration import control_database  # noqa: F401
from test_hr_panorama_database import _seed_owner_scope

from app.agent_brain.models import load_capability_cards
from app.control_plane.auth import AuthSecrets
from app.control_plane.models import AuthContext, IdentityMode, Role
from app.hr.panorama_analysis import PanoramaAnalyzer
from app.hr.panorama_collection import (
    CollectionError,
    CollectionResult,
    NormalizedPublicJob,
    SourceTarget,
)
from app.hr.panorama_context import PanoramaContextProvider
from app.hr.panorama_evidence import EvidenceArchive, EvidencePayload
from app.hr.panorama_models import CreateProductionBatch, CreateTalentSource
from app.hr.panorama_producer import PanoramaProductionPipeline
from app.hr.panorama_repository import PanoramaRepository
from app.hr.panorama_routes import build_panorama_router
from app.hr.panorama_service import PanoramaService

NOW = datetime(2026, 9, 6, 8, tzinfo=UTC)


class _AllowHrAgent:
    def decide_for_user_id(self, owner_id: UUID, agent_id: str):
        assert isinstance(owner_id, UUID)
        assert agent_id == "hr-bot"
        return SimpleNamespace(allowed=True)

    def permitted_agents_for_user_id(self, _owner_id: UUID):
        return tuple(
            card for card in load_capability_cards() if card.agent_id == "hr-bot"
        )

    def permitted_catalog_for_user_id(self, _owner_id: UUID):
        return ()


class _IdentityAuth:
    mode = IdentityMode.PRODUCTION
    route_prefix = "/"
    cookie_name = "panorama-session"
    csrf_cookie_name = "panorama-csrf"
    public_base_url = "https://agent.example.test"
    trusted_proxy_networks = ()
    rate_limiter = None
    secrets = AuthSecrets(b"s" * 32, key_version=1)

    def __init__(self, owner_id: UUID) -> None:
        self._context = AuthContext(owner_id, Role.PLATFORM_OWNER, uuid4(), False)

    def authenticate(self, token: str):
        return (self._context, "csrf") if token == "valid" else None

    @staticmethod
    def verify_csrf(submitted: str, expected: str) -> bool:
        return submitted == expected == "csrf"


class _AnalysisModel:
    version = "configured:gpt:acceptance-v1"

    async def generate_json(self, stage, payload):
        jobs = payload["jobs"]
        facts = [
            {
                "fact_id": f"{stage}-fact-{index}",
                "text": f"公开招聘{job['title']}，工作地点为{job['location']}",
                "snapshot_id": job["snapshot_id"],
                "observation_id": job["observation_id"],
                "source_url": job["source_url"],
                "observed_at": job["observed_at"],
            }
            for index, job in enumerate(jobs, start=1)
        ]
        return {
            "facts": facts,
            "inferences": [
                {
                    "text": "精密结构和量产能力是当前共同招聘信号",
                    "basis_fact_ids": [fact["fact_id"] for fact in facts],
                }
            ],
            "unknowns": [{"text": "实际招聘人数和产品项目归属未公开"}],
            "direction_clusters": {"结构": len(jobs)},
            "summary": "两家重点企业持续招聘精密结构与量产人才。",
        }


class _Collector:
    def __init__(self, archive: EvidenceArchive, failed_source_id: UUID) -> None:
        self._archive = archive
        self._failed_source_id = failed_source_id

    async def collect(self, target: SourceTarget) -> CollectionResult:
        if target.source_id == self._failed_source_id:
            raise CollectionError("source_unavailable")
        body = (
            f'{{"company":"{target.company_name}","jobs":'
            f'{{"title":"高级结构工程师"}}}}'
        ).encode()
        evidence = self._archive.store(
            EvidencePayload(
                source_url=target.source_url,
                mime="application/json",
                body=body,
                response_headers={"Content-Type": "application/json"},
            )
        )
        return CollectionResult(
            target=target,
            jobs=(
                NormalizedPublicJob(
                    public_job_key=f"structure-{target.source_id}",
                    title="高级结构工程师",
                    location="深圳",
                    duty_excerpt="负责精密结构研发、喷嘴和挤出系统量产",
                    requirement_excerpt="五年以上精密结构、可靠性和制造工艺经验",
                    source_url=f"{target.source_url}/positions/structure",
                ),
            ),
            evidence=evidence,
            observed_at=NOW,
        )


@pytest.mark.postgres
@pytest.mark.asyncio
async def test_background_intelligence_keeps_raw_jobs_and_ai_analysis_then_serves_both(
    control_database, tmp_path,  # noqa: F811
) -> None:
    environment = control_database["environments"]["production"]
    with psycopg.connect(environment["admin"]) as admin:
        scope = _seed_owner_scope(admin, "Panorama production acceptance")
    repository = PanoramaRepository(environment["urls"]["platform_control_app"])
    companies = ("联合光电", "禾赛科技", "拓竹科技")
    sources = tuple(
        repository.create_source(
            CreateTalentSource(
                source_id=uuid4(),
                owner_id=scope["owner"],
                client_request_id=uuid4(),
                company_key=f"acceptance-{index}-{uuid4().hex[:8]}",
                canonical_name=company,
                aliases=(),
                approved_urls=(f"https://company-{index}.example.com/jobs",),
                active=True,
            )
        )
        for index, company in enumerate(companies, start=1)
    )
    targets = tuple(
        SourceTarget(
            source.source_id,
            source.canonical_name,
            source.approved_urls[0],
            source.approved_urls,
        )
        for source in sources
    )
    batch = repository.create_production_batch(
        CreateProductionBatch(
            batch_id=uuid4(),
            owner_id=scope["owner"],
            client_request_id=uuid4(),
            selected_source_ids=tuple(source.source_id for source in sources),
            trigger_kind="schedule",
            analyzer_version=_AnalysisModel.version,
        )
    )
    archive = EvidenceArchive(tmp_path / "evidence")

    delivery = await PanoramaProductionPipeline(
        batch=batch,
        targets=targets,
        collector=_Collector(archive, sources[2].source_id),
        analyzer=PanoramaAnalyzer(_AnalysisModel()),
        repository=repository,
    ).run()

    assert delivery.coverage_state == "partial"
    assert delivery.snapshot_count == 2
    service = PanoramaService(repository, evidence_archive=archive)
    report = service.current_report()
    assert report is not None
    assert report.publication is not None
    assert report.publication.batch_id == batch.batch_id
    assert report.insight.model_version == _AnalysisModel.version
    assert len(report.snapshots) == 2
    assert len(report.evidence_attempts) == 3
    assert all(snapshot.duty_excerpt for snapshot in report.snapshots)
    assert all(fact["snapshot_id"] for fact in report.insight.facts)
    assert report.insight.inferences[0]["basis_fact_ids"]

    app = FastAPI()

    async def authorize(_request: Request, *, writable: bool = False):
        assert writable is False
        return scope["owner"]

    app.include_router(build_panorama_router(service, authorize))
    with TestClient(app) as client:
        current = client.get("/api/hr/panorama/current")
        assert current.status_code == 200
        payload = current.json()
        assert len(payload["snapshots"]) == 2
        assert payload["insight"]["inferences"]
        assert payload["publication"]["coverage_state"] == "partial"
        assert len(payload["evidence"]) == 2

        evidence = payload["evidence"][0]
        raw = client.get(
            f"/api/hr/panorama/reports/{payload['publication']['publication_id']}"
            f"/evidence/{evidence['sha256']}"
        )
        assert raw.status_code == 200
        assert raw.content == archive.read(evidence["sha256"])

        exported = client.get(
            f"/api/hr/panorama/reports/{payload['publication']['publication_id']}"
            "/export?format=xlsx"
        )
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        assert {"原始岗位", "AI分析", "来源覆盖", "证据索引"} <= set(
            workbook.sheetnames
        )

    fragment = PanoramaContextProvider(repository, now=lambda: NOW).for_turn(
        scope["owner"],
        scope["position"],
        "参考招聘全景分析，完善高级结构工程师人才画像",
        scope["turn"],
        task_kind="talent_profile",
    )
    assert fragment is not None
    assert fragment.publication_id == report.publication.publication_id
    assert fragment.insight_version_ids == (report.insight.insight_version_id,)
    assert fragment.facts
    assert fragment.inferences
    assert set(fragment.source_urls) == {
        snapshot.source_url for snapshot in report.snapshots
    }
