import base64
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from tools.sdk_q4_repair import (
    RepairConflict,
    assert_source_story,
    build_source_note_sql,
    change_story_body,
    parse_q4_workbook,
    render_story,
    select_rows,
)


def _workbook(path: Path) -> bytes:
    workbook = Workbook()
    q4 = workbook.active
    q4.title = "2026年Q4上线版本"
    q4.append(["平台", "任务", "责任人", "版本", "开始时间", "结束时间", "优先级", "风险点", "完成情况", "备注"])
    q4.append(["Android Wrapper", "根据客户需要更新版本", "千手修罗", None, None, None, "中"])
    q4.append(["OrbbecViewer 优化", "支持多相机连接与界面显示优化", "夜枫", "版本2",
               date(2026, 10, 1), date(2026, 10, 31), "高", "需兼容旧设备", "开发中", "设计文档"])
    q4["J3"].hyperlink = "https://example.com/design"
    q3 = workbook.create_sheet("2026年Q3上线版本")
    q3.append(["平台", "任务", "责任人"])
    q3.append(["Android Wrapper", "不应导入的Q3任务", "某人"])
    workbook.save(path)
    return path.read_bytes()


def test_parser_selects_exact_q4_sheet_and_preserves_rows_links(tmp_path):
    rows = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))
    assert [(r.excel_row, r.module, r.title) for r in rows] == [
        (2, "Android Wrapper", "根据客户需要更新版本"),
        (3, "OrbbecViewer 优化", "支持多相机连接与界面显示优化"),
    ]
    assert rows[1].links == (("J3", "设计文档", "https://example.com/design"),)
    assert rows[1].planned_start == "2026-10-01"
    assert rows[1].planned_end == "2026-10-31"


def test_renderer_marks_vague_android_row_without_fabricated_proposer(tmp_path):
    row = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))[0]
    rendered = render_story(row)
    assert rendered["sourceNote"] == "2026年Q4上线版本｜Android Wrapper｜Excel第2行"
    assert "根据客户需要更新版本" in rendered["spec"]
    assert "客户" in rendered["spec"] and "目标版本" in rendered["spec"]
    assert "验收标准" in rendered["verify"]
    assert "千手修罗" not in rendered["spec"]
    assert "SHA-256" not in rendered["spec"]
    assert "SDK2026PLAN" not in rendered["spec"]
    assert "提出人：千手修罗" not in rendered["spec"]


def test_renderer_keeps_source_dates_risk_progress_and_link(tmp_path):
    row = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))[1]
    rendered = render_story(row)
    for value in ("版本2", "2026-10-01", "2026-10-31", "需兼容旧设备", "开发中", "设计文档"):
        assert value in rendered["spec"]
    assert 'href="https://example.com/design"' in rendered["spec"]


def test_matching_accepts_zentao_auto_link_without_losing_reference(tmp_path):
    rows = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))
    url = "https://example.com/design"
    row = replace(rows[1], note=f"参考资料：{url}")
    rendered = render_story(row)
    actual_spec = rendered["spec"].replace(
        f"参考资料：{url}", f'参考资料：<a href="{url}" target="_blank">{url}</a>'
    )
    story = {
        "id": 13, "title": row.title, "product": 1, "module": 0, "pri": 1,
        "status": "reviewing", "reviewer": ["cangyuan"], "assignedTo": "sdk_owner",
        "openedBy": "sdk-api", "deleted": False, "sourceNote": "",
        "spec": actual_spec, "verify": rendered["verify"], "version": 2,
    }
    assert assert_source_story(story, row, product_id=1) == "body_done"


def test_matching_rejects_human_changes_and_accepts_idempotent_result(tmp_path):
    row = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))[0]
    story = {
        "id": 12, "title": row.title, "product": 1, "pri": 2, "status": "reviewing",
        "deleted": False, "sourceNote": "", "version": 1,
        "spec": ("2026 SDK 年度任务计划转入的待评审需求。"
                 "导入标识：SDK2026PLAN-41df018a-R0002；"
                 "源文件 SHA-256：41df018a029f1d27417bec3008bbe7aab3becd785cc38b312c563ce18ee8370d"),
    }
    assert_source_story(story, row, product_id=1)
    edited = {**story, "spec": "有人已经修改了需求描述"}
    with pytest.raises(RepairConflict):
        assert_source_story(edited, row, product_id=1)
    with pytest.raises(RepairConflict, match="human_edit"):
        assert_source_story({**story, "version": 2}, row, product_id=1)
    repaired = {**story, **render_story(row)}
    assert_source_story(repaired, row, product_id=1)
    source_done = {**story, "sourceNote": render_story(row)["sourceNote"]}
    assert assert_source_story(source_done, row, product_id=1) == "source_done"
    body_done = {**story, "spec": render_story(row)["spec"],
                 "verify": render_story(row)["verify"], "version": 2}
    assert assert_source_story(body_done, row, product_id=1) == "body_done"


def test_parser_requires_q4_sheet_and_rejects_wrong_header(tmp_path):
    workbook = Workbook()
    workbook.active.title = "2026年Q3上线版本"
    path = tmp_path / "wrong.xlsx"
    workbook.save(path)
    with pytest.raises(RepairConflict, match="q4_sheet_missing"):
        parse_q4_workbook(path.read_bytes())
    workbook.active.title = "2026年Q4上线版本"
    workbook.active.append(["bad", "header"])
    workbook.save(path)
    with pytest.raises(RepairConflict, match="q4_header_invalid"):
        parse_q4_workbook(path.read_bytes())


def test_canary_selection_requires_exact_excel_row(tmp_path):
    rows = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))
    assert select_rows(rows, only_row=3, limit=None) == (rows[1],)
    with pytest.raises(RepairConflict, match="q4_row_not_found"):
        select_rows(rows, only_row=92, limit=None)


def test_source_note_sql_is_guarded_and_excludes_already_updated_row(tmp_path):
    rows = parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))
    stories = {}
    for row in rows:
        rendered = render_story(row)
        stories[row.excel_row] = {
            "id": row.excel_row + 10, "title": row.title, "product": 1,
            "module": 0, "pri": 2 if row.excel_row == 2 else 1,
            "status": "reviewing", "reviewer": ["cangyuan"],
            "assignedTo": "sdk_owner", "openedBy": "sdk-api", "deleted": False,
            "sourceNote": rendered["sourceNote"] if row.excel_row == 2 else "",
            "spec": rendered["spec"], "verify": rendered["verify"], "version": 2,
        }
    sql = build_source_note_sql(rows, stories, product_id=1)
    assert "WHERE id=13 AND product=1 AND status='reviewing' AND version=2" in sql
    assert "WHERE id=12" not in sql
    assert "IF ROW_COUNT() <> 1" in sql
    assert "ROLLBACK" in sql and "COMMIT" in sql
    assert base64.b64encode(render_story(rows[1])["sourceNote"].encode()).decode() in sql


@pytest.mark.asyncio
async def test_change_uses_native_change_endpoint_and_preserves_assignment(tmp_path):
    row = replace(parse_q4_workbook(_workbook(tmp_path / "plan.xlsx"))[0], title="A&B")
    before = {
        "id": 12, "title": "A&amp;B", "product": 1, "module": 0, "pri": 2,
        "status": "reviewing", "reviewer": ["cangyuan"], "assignedTo": "sdk_owner",
        "openedBy": "sdk-api", "deleted": False, "sourceNote": "",
        "spec": ("2026 SDK 年度任务计划转入的待评审需求。"
                 "导入标识：SDK2026PLAN-41df018a-R0002；"
                 "源文件 SHA-256：41df018a029f1d27417bec3008bbe7aab3becd785cc38b312c563ce18ee8370d"),
        "verify": "待补全", "version": 1, "lastEditedDate": "2026-10-20 10:00:00",
    }

    class FakeClient:
        async def _request(self, method, endpoint, *, payload, api_version, expect_envelope):
            assert (method, endpoint, api_version, expect_envelope) == ("POST", "stories/12/change", 1, False)
            assert payload["title"] == "A&B"
            assert payload["reviewer"] == ["cangyuan"]
            assert payload["lastEditedDate"] == before["lastEditedDate"]
            return {"id": 12}

        async def get_story(self, story_id):
            return {**before, **render_story(row), "sourceNote": "", "version": 2}

    after = await change_story_body(FakeClient(), before, row, product_id=1)
    assert after["assignedTo"] == "sdk_owner"
    assert assert_source_story(after, row, product_id=1) == "body_done"
